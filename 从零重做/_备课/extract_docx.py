# -*- coding: utf-8 -*-
"""把实验指导书 .docx 提取成纯文本，供备课参考。

用法: python extract_docx.py <输入.docx> <输出.txt>
"""
import sys

from docx import Document

src, out = sys.argv[1], sys.argv[2]
doc = Document(src)

lines = []
for para in doc.paragraphs:
    t = para.text.strip()
    if t:
        lines.append(t)

for i, table in enumerate(doc.tables, 1):
    lines.append(f"=== 表格 {i} ===")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        lines.append(" | ".join(cells))

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"OK: {len(lines)} lines -> {out}")